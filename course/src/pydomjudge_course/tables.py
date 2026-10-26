"""Tables of exports (Moodle and others) as rows of strings: CSV, TSV and xlsx.

``read_table`` detects the encoding (with or without BOM), the delimiter of
text files and reads xlsx sheets; every cell becomes a string. ``write_table``
writes a table back with its header order, encoding, delimiter and BOM, so that
an export changed in a few columns can be imported again.

xlsx stores numbers, not text: an integral number becomes ``"123"``, so leading
zeros of identifiers are lost (``"00123"`` reads as ``"123"``). Where
identifiers can have leading zeros, the roster has to declare their padding.

``python -m pydomjudge_course.tables FILE...`` prints what ``probe`` finds
(format, encoding, delimiter, BOM, column names, row count) and never a cell
value, to check the format of real exports without showing their contents.
For JSON files it prints the shape: lists and their lengths, and for records
(objects in a list) the union of their keys with the types of the values.
Keys of other objects are not shown, as they may be data (logins, names).
"""

import argparse
import codecs
import csv
import dataclasses
import datetime
import io
import json
import pathlib
import re
import sys
import zipfile
from collections.abc import Iterable, Sequence
from typing import Literal

Format = Literal["csv", "xlsx"]

TEXT_DELIMITERS = ",;\t"


class TableError(ValueError):
    """A file that is no table this module can read; the message names the file
    and positions, never cell contents."""


@dataclasses.dataclass
class Table:
    header: tuple[str, ...]
    rows: list[dict[str, str]] = dataclasses.field(default_factory=list)
    # How the file was read, and how write_table writes it by default
    format: Format = "csv"
    delimiter: str = ","
    encoding: str = "utf-8"
    bom: bool = False
    source: pathlib.Path | None = None

    def column(self, *candidates: str, required: bool = True) -> str | None:
        """The header name matching one of the candidates, ignoring case,
        spaces, underscores and hyphens ("Email address" = "emailaddress")."""
        by_key = {_column_key(name): name for name in self.header}
        for candidate in candidates:
            if (name := by_key.get(_column_key(candidate))) is not None:
                return name
        if required:
            raise TableError(
                f"{self._where()}: no column {' or '.join(map(repr, candidates))} "
                f"(columns: {', '.join(self.header)})"
            )
        return None

    def columns_starting(self, prefix: str) -> list[str]:
        """The header names starting with prefix (case-insensitive), in order."""
        prefix = prefix.casefold()
        return [name for name in self.header if name.casefold().startswith(prefix)]

    def _where(self) -> str:
        return str(self.source) if self.source is not None else "table"


def _column_key(name: str) -> str:
    return re.sub(r"[\s_-]+", "", name).casefold()


def read_table(
    path: pathlib.Path, *, sheet: str | None = None, encoding: str | None = None
) -> Table:
    """The table in path: xlsx by suffix, otherwise delimited text. Without
    encoding, text files are read as UTF-8 (with or without BOM) or, failing
    that, in the detected encoding."""
    if path.suffix.lower() == ".xlsx":
        if encoding is not None:
            raise TableError(f"{path}: xlsx files have no text encoding")
        return _read_xlsx(path, sheet)
    if sheet is not None:
        raise TableError(f"{path}: only xlsx files have sheets")
    return _read_text(path, encoding)


def write_table(
    table: Table,
    path: pathlib.Path,
    *,
    encoding: str | None = None,
    bom: bool | None = None,
    delimiter: str | None = None,
) -> None:
    """Write table as delimited text, by default as it was read (header order,
    encoding, delimiter, BOM). Rows may lack columns (written empty), but not
    have columns outside the header."""
    if path.suffix.lower() == ".xlsx":
        raise NotImplementedError("writing xlsx files; write CSV instead")
    if delimiter is None:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else table.delimiter
    encoding = encoding or table.encoding
    bom = table.bom if bom is None else bom
    columns = set(table.header)
    for number, row in enumerate(table.rows, start=1):
        if unknown := [name for name in row if name not in columns]:
            raise TableError(
                f"row {number} has columns outside the header: {', '.join(unknown)}"
            )
    with path.open("w", encoding=encoding, newline="") as f:
        if bom:
            f.write("\ufeff")
        writer = csv.writer(f, delimiter=delimiter)
        writer.writerow(table.header)
        for row in table.rows:
            writer.writerow([row.get(name, "") for name in table.header])


# Reading text files


# Among equally plausible single-byte encodings (charset_normalizer cannot tell
# cp1250 from cp1252 by German or French text), the one of Western-European
# Windows, which exports from there use
PREFERRED_ENCODING = "cp1252"


def _decode(
    path: pathlib.Path, data: bytes, encoding: str | None
) -> tuple[str, str, bool]:
    """Text, encoding and whether there was a UTF-8 BOM."""
    if encoding is not None:
        try:
            utf_8 = codecs.lookup(encoding).name in ("utf-8", "utf-8-sig")
        except LookupError as e:
            raise TableError(f"{path}: unknown encoding {encoding}") from e
        if utf_8 and data.startswith(codecs.BOM_UTF8):
            data = data[len(codecs.BOM_UTF8) :]
            bom = True
        else:
            bom = False
        try:
            return data.decode("utf-8" if utf_8 else encoding), encoding, bom
        except UnicodeDecodeError as e:
            raise TableError(f"{path}: not {encoding} ({e.__class__.__name__})") from e
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8) :].decode("utf-8"), "utf-8", True
    try:
        return data.decode("utf-8"), "utf-8", False
    except UnicodeDecodeError:
        pass
    import charset_normalizer

    matches = charset_normalizer.from_bytes(data)
    best = matches.best()
    if best is None:
        raise TableError(f"{path}: cannot detect the text encoding")
    tied = [
        match
        for match in matches
        if match.chaos == best.chaos and match.coherence == best.coherence
    ]
    chosen = next(
        (match for match in tied if match.encoding == PREFERRED_ENCODING), best
    )
    return str(chosen), chosen.encoding, False


def _sniff_delimiter(path: pathlib.Path, text: str) -> str:
    if path.suffix.lower() == ".tsv":
        return "\t"
    sample = "\n".join(text.splitlines()[:20])
    try:
        return csv.Sniffer().sniff(sample, delimiters=TEXT_DELIMITERS).delimiter
    except csv.Error:
        # One column, or nothing the sniffer recognises
        return ","


def _read_text(path: pathlib.Path, encoding: str | None) -> Table:
    try:
        data = path.read_bytes()
    except OSError as e:
        raise TableError(f"{path}: {e.strerror}") from e
    text, encoding, bom = _decode(path, data, encoding)
    delimiter = _sniff_delimiter(path, text)
    records = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    header, rows = _rows_of(path, records)
    return Table(
        header=header,
        rows=rows,
        format="csv",
        delimiter=delimiter,
        encoding=encoding,
        bom=bom,
        source=path,
    )


# Reading xlsx files


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime | datetime.date | datetime.time):
        return value.isoformat()
    return str(value)


def _read_xlsx(path: pathlib.Path, sheet: str | None) -> Table:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as e:
        raise TableError(f"{path}: not a readable xlsx file ({e})") from e
    try:
        if sheet is not None:
            if sheet not in workbook.sheetnames:
                raise TableError(
                    f"{path}: no sheet {sheet!r} "
                    f"(sheets: {', '.join(workbook.sheetnames)})"
                )
            worksheet = workbook[sheet]
        elif len(workbook.sheetnames) != 1:
            raise TableError(
                f"{path}: {len(workbook.sheetnames)} sheets "
                f"({', '.join(workbook.sheetnames)}); name the one to read"
            )
        else:
            worksheet = workbook[workbook.sheetnames[0]]
        records = [
            [_cell_text(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()
    header, rows = _rows_of(path, records)
    return Table(header=header, rows=rows, format="xlsx", source=path)


# Shared


def _rows_of(
    path: pathlib.Path, records: Sequence[Sequence[str]]
) -> tuple[tuple[str, ...], list[dict[str, str]]]:
    """Header and rows of records: rows without any value are skipped, short
    rows padded, empty cells after the last column (a trailing delimiter, or
    empty xlsx cells) tolerated."""
    records = [record for record in records if any(cell.strip() for cell in record)]
    if not records:
        raise TableError(f"{path}: no header")
    header = list(records[0])
    while header and not header[-1].strip():
        header.pop()
    width = len(header)
    seen: set[str] = set()
    for position, name in enumerate(header, start=1):
        if not name.strip():
            raise TableError(f"{path}: column {position} has no name")
        if name in seen:
            raise TableError(f"{path}: column {name!r} appears twice")
        seen.add(name)
    rows = []
    for number, record in enumerate(records[1:], start=2):
        if any(cell.strip() for cell in record[width:]):
            raise TableError(
                f"{path}: row {number} has values after the last column "
                f"({len(record)} cells, {width} columns)"
            )
        cells = list(record[:width]) + [""] * (width - len(record))
        rows.append(dict(zip(header, cells, strict=True)))
    return tuple(header), rows


# Probe


@dataclasses.dataclass(frozen=True)
class Probe:
    """What a file looks like, without its contents."""

    path: pathlib.Path
    format: Format
    encoding: str
    delimiter: str
    bom: bool
    header: tuple[str, ...]
    row_count: int

    def describe(self) -> str:
        delimiter = {"\t": "tab", ",": "comma", ";": "semicolon"}.get(
            self.delimiter, repr(self.delimiter)
        )
        layout = (
            "xlsx"
            if self.format == "xlsx"
            else f"text, {self.encoding}{' with BOM' if self.bom else ''}, "
            f"{delimiter}-separated"
        )
        columns = "\n".join(f"  {name}" for name in self.header)
        return (
            f"{self.path}: {layout}, {len(self.header)} columns, "
            f"{self.row_count} rows\n{columns}"
        )


@dataclasses.dataclass(frozen=True)
class JsonProbe:
    """The shape of a JSON file, without its values."""

    path: pathlib.Path
    encoding: str
    bom: bool
    shape: str

    def describe(self) -> str:
        return (
            f"{self.path}: JSON, {self.encoding}{' with BOM' if self.bom else ''}\n"
            f"  {self.shape}"
        )


def _json_type(value: object) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int | float):
        return "number"
    if isinstance(value, str):
        return "string"
    if isinstance(value, list):
        return "list"
    return "object"


def _json_shape(value: object) -> str:
    if isinstance(value, list):
        return f"list of {len(value)}" + (f": {_json_items(value)}" if value else "")
    if isinstance(value, dict):
        # Not records, so the keys may be data: only their number
        return f"object with {len(value)} keys"
    return _json_type(value)


def _json_items(items: list) -> str:
    """The merged shape of the elements of a list."""
    if all(isinstance(item, dict) for item in items):
        types: dict[str, set[str]] = {}
        for item in items:
            for key, value in item.items():
                types.setdefault(key, set()).add(_json_type(value))
        keys = ", ".join(
            f"{key} ({'|'.join(sorted(types[key]))})" for key in sorted(types)
        )
        return f"objects with keys {keys}"
    if all(isinstance(item, list) for item in items):
        lengths = sorted({len(item) for item in items})
        length = (
            str(lengths[0]) if len(lengths) == 1 else f"{lengths[0]} to {lengths[-1]}"
        )
        inner = [element for item in items for element in item]
        return f"lists of {length}" + (f": {_json_items(inner)}" if inner else "")
    return "values of type " + "|".join(sorted({_json_type(item) for item in items}))


def probe_json(path: pathlib.Path) -> JsonProbe:
    try:
        data = path.read_bytes()
    except OSError as e:
        raise TableError(f"{path}: {e.strerror}") from e
    text, encoding, bom = _decode(path, data, None)
    try:
        value = json.loads(text)
    except json.JSONDecodeError as e:
        # The message names the position only
        raise TableError(
            f"{path}: invalid JSON at line {e.lineno} column {e.colno}"
        ) from None
    return JsonProbe(path=path, encoding=encoding, bom=bom, shape=_json_shape(value))


def probe(path: pathlib.Path) -> Probe | JsonProbe:
    """What a file looks like (a table, or the shape of a JSON file)."""
    if path.suffix.lower() == ".json":
        return probe_json(path)
    table = read_table(path)
    return Probe(
        path=path,
        format=table.format,
        encoding=table.encoding,
        delimiter=table.delimiter,
        bom=table.bom,
        header=table.header,
        row_count=len(table.rows),
    )


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m pydomjudge_course.tables",
        description="Show format and column names of tables (or the shape of "
        "JSON files), never their values",
    )
    parser.add_argument("files", nargs="+", type=pathlib.Path)
    args = parser.parse_args(list(argv) if argv is not None else None)
    status = 0
    for path in args.files:
        try:
            print(probe(path).describe())
        except TableError as e:
            print(f"error: {e}", file=sys.stderr)
            status = 1
    return status


if __name__ == "__main__":
    sys.exit(main())

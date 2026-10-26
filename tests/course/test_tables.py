"""Reading and writing tables, on synthetic files only."""

import datetime
import pathlib

import openpyxl
import pytest

from pydomjudge_course import tables
from pydomjudge_course.tables import Table, TableError, read_table, write_table

HEADER = ["Identifier", "Full name", "Email address", "Grade", "Feedback comments"]
ROWS = [
    ["Participant 1", "Jürgen Müller", "j.mueller@example.org", "", ""],
    ["Participant 2", "Zoë Ångström", "z.angstrom@example.org", "1", "Line 1\nLine 2"],
]


def _csv_text(delimiter: str = ",") -> str:
    import csv
    import io

    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter)
    writer.writerow(HEADER)
    writer.writerows(ROWS)
    return buffer.getvalue()


def _expected_rows() -> list[dict[str, str]]:
    return [dict(zip(HEADER, row, strict=True)) for row in ROWS]


@pytest.mark.parametrize(
    ("name", "encoding", "bom", "delimiter"),
    [
        ("plain.csv", "utf-8", False, ","),
        ("bom.csv", "utf-8", True, ","),
        ("semicolon.csv", "utf-8", False, ";"),
        ("choice.txt", "utf-8", False, "\t"),
        ("sheet.tsv", "utf-8", False, "\t"),
        ("windows.csv", "cp1252", False, ","),
    ],
)
def test_read_text(tmp_path: pathlib.Path, name, encoding, bom, delimiter):
    path = tmp_path / name
    text = _csv_text(delimiter)
    if encoding == "cp1252":
        # Enough non-ASCII text for the encoding detection
        text += "Participant 3,André Désirée Çelik,a@example.org,,Grüße\n" * 5
    data = text.encode(encoding)
    path.write_bytes((b"\xef\xbb\xbf" if bom else b"") + data)

    table = read_table(path)
    assert table.header == tuple(HEADER)
    assert table.rows[:2] == _expected_rows()
    assert (table.delimiter, table.bom, table.format) == (delimiter, bom, "csv")
    if encoding == "cp1252":
        assert table.rows[2]["Full name"] == "André Désirée Çelik"
        assert table.encoding == "cp1252"
    else:
        assert table.encoding == "utf-8"


def test_round_trip_keeps_header_order_bom_and_cells(tmp_path: pathlib.Path):
    source = tmp_path / "worksheet.csv"
    source.write_bytes(b"\xef\xbb\xbf" + _csv_text().encode())
    table = read_table(source)
    table.rows[0]["Grade"] = "1"
    target = tmp_path / "updated.csv"
    write_table(table, target)

    again = read_table(target)
    assert again.header == tuple(HEADER)
    assert again.bom
    assert again.rows[0]["Grade"] == "1"
    assert again.rows[1] == _expected_rows()[1]
    assert target.read_bytes().startswith(b"\xef\xbb\xbfIdentifier,Full name")


def test_write_new_table_and_options(tmp_path: pathlib.Path):
    table = Table(header=("a", "b"), rows=[{"a": "1"}, {"a": "2", "b": "x"}])
    write_table(table, tmp_path / "t.csv", delimiter=";")
    assert (tmp_path / "t.csv").read_bytes() == b"a;b\r\n1;\r\n2;x\r\n"
    write_table(table, tmp_path / "t.tsv")
    assert read_table(tmp_path / "t.tsv").rows == [
        {"a": "1", "b": ""},
        {"a": "2", "b": "x"},
    ]
    with pytest.raises(TableError, match="outside the header: c"):
        write_table(Table(header=("a",), rows=[{"c": "1"}]), tmp_path / "u.csv")
    with pytest.raises(NotImplementedError):
        write_table(table, tmp_path / "t.xlsx")


def test_lenient_and_strict_rows(tmp_path: pathlib.Path):
    path = tmp_path / "rows.csv"
    # trailing delimiters, a short row, an empty row
    path.write_text("a,b,\n1,2,\n3\n,,\n\n4,5,,\n")
    table = read_table(path)
    assert table.header == ("a", "b")
    assert table.rows == [
        {"a": "1", "b": "2"},
        {"a": "3", "b": ""},
        {"a": "4", "b": "5"},
    ]

    path.write_text("a,b\n1,2,surprise\n")
    with pytest.raises(TableError, match="row 2 has values after the last column"):
        read_table(path)


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("a,b,a\n1,2,3\n", "column 'a' appears twice"),
        ("a,,b\n1,2,3\n", "column 2 has no name"),
        ("\n\n", "no header"),
    ],
)
def test_bad_headers(tmp_path: pathlib.Path, content, message):
    path = tmp_path / "bad.csv"
    path.write_text(content)
    with pytest.raises(TableError, match=message):
        read_table(path)


def _workbook(path: pathlib.Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)


def test_read_xlsx_cells_are_strings(tmp_path: pathlib.Path):
    path = tmp_path / "grades.xlsx"
    _workbook(
        path,
        {
            "Sheet": [
                ["Student ID", "Grade", "Ratio", "Due", "Done", "Note", None],
                [123, 100.0, 0.25, datetime.datetime(2030, 1, 2, 3, 4), True, None],
                [None, None, None, None, None, None, None],
                ["00123", 7, 1.5, datetime.date(2030, 1, 2), False, "text"],
            ]
        },
    )
    table = read_table(path)
    assert table.format == "xlsx"
    assert table.header == ("Student ID", "Grade", "Ratio", "Due", "Done", "Note")
    assert table.rows == [
        {
            "Student ID": "123",
            "Grade": "100",
            "Ratio": "0.25",
            "Due": "2030-01-02T03:04:00",
            "Done": "TRUE",
            "Note": "",
        },
        {
            "Student ID": "00123",  # stored as text, so the zeros stay
            "Grade": "7",
            "Ratio": "1.5",
            "Due": "2030-01-02T00:00:00",
            "Done": "FALSE",
            "Note": "text",
        },
    ]


def test_xlsx_sheets(tmp_path: pathlib.Path):
    path = tmp_path / "two.xlsx"
    _workbook(path, {"First": [["a"], ["1"]], "Second": [["b"], ["2"]]})
    with pytest.raises(TableError, match="2 sheets .*First, Second"):
        read_table(path)
    assert read_table(path, sheet="Second").rows == [{"b": "2"}]
    with pytest.raises(TableError, match="no sheet 'Third'"):
        read_table(path, sheet="Third")
    (tmp_path / "broken.xlsx").write_bytes(b"not a zip")
    with pytest.raises(TableError, match="not a readable xlsx file"):
        read_table(tmp_path / "broken.xlsx")


def test_column_lookup():
    table = Table(
        header=("Email address", "Student ID", "Grade can be changed", "Grade")
    )
    assert table.column("emailaddress") == "Email address"
    assert table.column("student_id", "Username") == "Student ID"
    assert table.column("Username", required=False) is None
    with pytest.raises(TableError, match="no column 'Username'"):
        table.column("Username")
    assert table.columns_starting("grade") == ["Grade can be changed", "Grade"]


def test_probe_never_prints_cells(tmp_path: pathlib.Path, capsys):
    secret = "SECRET-CELL-4711"
    good = tmp_path / "export.csv"
    good.write_text(f"Name,Mail\n{secret},{secret}@example.org\n")
    workbook = tmp_path / "export.xlsx"
    _workbook(workbook, {"S": [["Name", "Mail"], [secret, 4711]]})
    duplicate = tmp_path / "dup.csv"
    duplicate.write_text(f"Name,Name\n{secret},{secret}\n")
    extra = tmp_path / "extra.csv"
    extra.write_text(f"Name\n{secret},{secret}\n")

    status = tables.main([str(good), str(workbook), str(duplicate), str(extra)])
    output = capsys.readouterr()
    assert status == 1
    assert secret not in output.out + output.err
    assert "4711" not in output.out + output.err
    assert "export.csv: text, utf-8, comma-separated, 2 columns, 1 rows" in output.out
    assert "  Name\n  Mail" in output.out
    assert "export.xlsx: xlsx, 2 columns, 1 rows" in output.out
    assert "column 'Name' appears twice" in output.err


def test_explicit_encoding(tmp_path: pathlib.Path):
    path = tmp_path / "latin.csv"
    path.write_bytes("Name\nŁukasz Żółć\n".encode("cp1250"))
    assert read_table(path, encoding="cp1250").rows == [{"Name": "Łukasz Żółć"}]
    path.write_bytes(b"Name\n\xff\xfe\n")
    with pytest.raises(TableError, match="not utf-8"):
        read_table(path, encoding="utf-8")


@pytest.mark.parametrize("encoding", ["utf-8", "UTF8", "utf-8-sig"])
def test_explicit_utf8_strips_bom(tmp_path: pathlib.Path, encoding):
    path = tmp_path / "bom.csv"
    path.write_bytes(b"\xef\xbb\xbfIdentifier,Grade\nParticipant 1,1\n")
    table = read_table(path, encoding=encoding)
    assert table.header == ("Identifier", "Grade")
    assert table.bom
    with pytest.raises(TableError, match="unknown encoding"):
        read_table(path, encoding="no-such-encoding")
